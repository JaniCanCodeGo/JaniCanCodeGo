#!/usr/bin/env python3
"""
CRR Civil Rights Review Document Processing Agent

Reads a completed CRR Summary of Findings (LOF), extracts Required Corrective
Actions, fills the Voluntary Compliance Plan (VCP), and generates the appropriate
LOF Cover Letter (Findings or No Findings).

Setup:
  pip install python-docx requests beautifulsoup4

  Place the three blank template .docx files in a templates/ folder:
    templates/VCP_template.docx
    templates/LOF_Findings_template.docx
    templates/LOF_NoFindings_template.docx

Usage:
  python3 crr_agent.py path/to/Summary_of_Findings.docx [--output-dir outputs/]
"""

import argparse
import re
import sys
from datetime import datetime, timedelta
from pathlib import Path

from docx import Document

SCRIPT_DIR = Path(__file__).parent
TEMPLATE_DIR = SCRIPT_DIR / "templates"

# California county codes (first 2 digits of CDS code)
CA_COUNTY_CODES = {
    "01": "Alameda",       "02": "Alpine",        "03": "Amador",
    "04": "Butte",         "05": "Calaveras",     "06": "Colusa",
    "07": "Contra Costa",  "08": "Del Norte",     "09": "El Dorado",
    "10": "Fresno",        "11": "Glenn",         "12": "Humboldt",
    "13": "Imperial",      "14": "Inyo",          "15": "Kern",
    "16": "Kings",         "17": "Lake",          "18": "Lassen",
    "19": "Los Angeles",   "20": "Madera",        "21": "Marin",
    "22": "Mariposa",      "23": "Mendocino",     "24": "Merced",
    "25": "Modoc",         "26": "Mono",          "27": "Monterey",
    "28": "Napa",          "29": "Nevada",        "30": "Orange",
    "31": "Placer",        "32": "Plumas",        "33": "Riverside",
    "34": "Sacramento",    "35": "San Benito",    "36": "San Bernardino",
    "37": "San Diego",     "38": "San Francisco", "39": "San Joaquin",
    "40": "San Luis Obispo", "41": "San Mateo",   "42": "Santa Barbara",
    "43": "Santa Clara",   "44": "Santa Cruz",    "45": "Shasta",
    "46": "Sierra",        "47": "Siskiyou",      "48": "Solano",
    "49": "Sonoma",        "50": "Stanislaus",    "51": "Sutter",
    "52": "Tehama",        "53": "Trinity",       "54": "Tulare",
    "55": "Tuolumne",      "56": "Ventura",       "57": "Yolo",
    "58": "Yuba",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def is_none_action(text: str) -> bool:
    """Return True if the text is a placeholder meaning 'no action required'."""
    t = text.strip().lower().rstrip(".")
    return t in ("none", "", "n/a", " " * 5)


def county_from_cds(cds_code: str) -> str:
    """Return the county name from the first 2 digits of a CDS code."""
    cds_clean = re.sub(r"[\s\-]", "", cds_code)
    return CA_COUNTY_CODES.get(cds_clean[:2], "")


def get_initials(name: str) -> str:
    """'Jessica Franklin' -> 'JF'"""
    return "".join(p[0].upper() for p in name.strip().split() if p)


def parse_principal(coordinator_str: str):
    """
    Parse 'Eric Preston, Principal' into (full_name, last_name, title).
    Handles 'First Last, Title' or just 'First Last'.
    """
    parts = coordinator_str.split(",", 1)
    full_name = parts[0].strip()
    title = parts[1].strip() if len(parts) > 1 else ""
    name_parts = full_name.split()
    last_name = name_parts[-1] if name_parts else ""
    return full_name, last_name, title


def normalize_cds(cds_code: str) -> str:
    """Normalize CDS code to 14 contiguous digits."""
    digits = re.sub(r"[\s\-]", "", cds_code)
    return digits.ljust(14, "0")[:14]


# ---------------------------------------------------------------------------
# CDE web helpers
# ---------------------------------------------------------------------------

_CDE_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept": "text/html,application/xhtml+xml",
}


def _fetch_cde_html(url: str, label: str = "") -> str:
    """Fetch a CDE page and return the HTML string, or '' on failure."""
    import urllib.request

    try:
        req = urllib.request.Request(url, headers=_CDE_HEADERS)
        with urllib.request.urlopen(req, timeout=15) as resp:
            html = resp.read().decode("utf-8", errors="ignore")
    except Exception as e:
        print(f"  [Info] CDE fetch failed{' (' + label + ')' if label else ''} "
              f"({type(e).__name__})")
        return ""

    if not html or "Host not in allowlist" in html or len(html) < 200:
        print(f"  [Info] CDE not reachable{' (' + label + ')' if label else ''}.")
        return ""

    return html


def _soup_label_value(html: str, label_lower: str) -> str:
    """
    Search an HTML page for a table row or dt/dd whose label contains
    label_lower, and return the adjacent value cell text.
    Falls back to regex when BeautifulSoup is not installed.
    """
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, "html.parser")

        for row in soup.find_all("tr"):
            cells = row.find_all(["th", "td"])
            if len(cells) >= 2:
                lbl = cells[0].get_text(" ", strip=True).lower()
                val = cells[1].get_text(" ", strip=True)
                if label_lower in lbl and val and val.lower() not in ("n/a", "none", ""):
                    return val

        for dt in soup.find_all("dt"):
            lbl = dt.get_text(strip=True).lower()
            dd = dt.find_next_sibling("dd")
            if dd and label_lower in lbl:
                val = dd.get_text(strip=True)
                if val and val.lower() not in ("n/a", "none", ""):
                    return val

    except ImportError:
        # Plain-text fallback: look for label followed by value on same or next line
        pattern = rf"{re.escape(label_lower)}[^:\n]*:?\s*([^\n<]{{3,80}})"
        m = re.search(pattern, html, re.IGNORECASE)
        if m:
            return m.group(1).strip()

    return ""


def _soup_all_label_values(html: str) -> dict:
    """Return all label→value pairs from a CDE details page as a dict."""
    result = {}
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, "html.parser")
        for row in soup.find_all("tr"):
            cells = row.find_all(["th", "td"])
            if len(cells) >= 2:
                lbl = cells[0].get_text(" ", strip=True).lower().strip()
                val = cells[1].get_text(" ", strip=True).strip()
                if lbl and val and val.lower() not in ("n/a", "none", ""):
                    result[lbl] = val
        for dt in soup.find_all("dt"):
            lbl = dt.get_text(strip=True).lower().strip()
            dd = dt.find_next_sibling("dd")
            if dd:
                val = dd.get_text(strip=True)
                if lbl and val and val.lower() not in ("n/a", "none", ""):
                    result[lbl] = val
    except ImportError:
        pass
    return result


# ---------------------------------------------------------------------------
# CDE School Directory lookups
# ---------------------------------------------------------------------------

def _parse_cde_detail_page(html: str) -> dict:
    """
    Extract school address, contact, and principal from a CDE school detail page.

    CDE detail pages use these exact field labels:
      School Address  => "123 Main St  City, CA 90000-1234  Google Map Link..."
      Mailing Address => "123 Main St  City, CA 90000-1234"
      Phone Number    => "(555) 123-4567"
      Administrator   => "Jane Smith Principal (555) 123-4567 Ext. 1 jsmith@school.edu"
      Email           => "jsmith@school.edu"  (sometimes "Information Not Available")
      County          => "Sacramento"
    """
    fields = _soup_all_label_values(html)

    def _pick(*keys):
        for k in keys:
            for fk, fv in fields.items():
                if k in fk.lower():
                    return fv
        return ""

    info = {}

    # --- Address: CDE puts full address in one field ---
    raw_addr = _pick("school address", "mailing address")
    if raw_addr:
        # Strip "Google Map Link..." trailer
        raw_addr = re.sub(r"\s*Google Map.*", "", raw_addr, flags=re.IGNORECASE).strip()

        # CDE format is typically: "123 Main St. Cityname, CA 90000"
        # or "123 Main St. City Name CA 90000" (no comma before CA)
        # Strategy: anchor on ", CA XXXXX" or " CA XXXXX" at the end,
        # then split street vs city at the last street-type suffix word.
        zip_m = re.search(r',?\s*CA\s+(\d{5})(?:-\d{4})?', raw_addr, re.IGNORECASE)
        if zip_m:
            zip_code  = zip_m.group(1)
            before_ca = raw_addr[:zip_m.start()].strip()

            # Identify the last street suffix — city follows it
            _SFX = (
                r"(?:Street|St|Avenue|Ave|Boulevard|Blvd|Road|Rd|Drive|Dr|"
                r"Way|Lane|Ln|Court|Ct|Circle|Cir|Place|Pl|Terrace|Ter|"
                r"Highway|Hwy|Parkway|Pkwy|Trail|Trl|Loop|Run|Row|Walk)\.?"
            )
            sfx_m = re.search(rf"(?i)\b({_SFX})\s+(.+)$", before_ca)
            if sfx_m:
                info["street"] = before_ca[: sfx_m.end(1)].strip().rstrip(",")
                info["city"]   = sfx_m.group(2).strip().rstrip(",")
            else:
                # Fallback: split on the last comma
                comma_idx = before_ca.rfind(",")
                if comma_idx > 0:
                    info["street"] = before_ca[:comma_idx].strip()
                    info["city"]   = before_ca[comma_idx + 1:].strip()
                else:
                    info["street"] = before_ca

            info["state"] = "CA"
            info["zip"]   = zip_code
        else:
            info["street"] = raw_addr  # store as-is if no CA zip found

    # --- Phone ---
    phone = _pick("phone number", "phone")
    if phone and phone.lower() not in ("information not available",):
        info["phone"] = phone

    # --- Administrator: name is first two words, email may be embedded ---
    admin = _pick("administrator")
    if admin:
        name_m = re.match(r"([A-Z][a-z]+(?:\s+[A-Z][a-z]+)+)", admin)
        if name_m:
            info["principal_name"] = name_m.group(1)
        emails = re.findall(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", admin)
        if emails:
            info["email"] = emails[0]

    # --- Email field fallback ---
    if not info.get("email"):
        email_val = _pick("email")
        if email_val and email_val.lower() not in ("information not available", "n/a", "none"):
            info["email"] = email_val

    # --- County (useful when looking up by name) ---
    county_val = _pick("county")
    if county_val:
        info["county"] = county_val

    # City/zip regex fallback if address field was missing
    if not info.get("street"):
        m = re.search(r"([A-Za-z ]{3,30}),\s*CA\s*(\d{5})", html)
        if m:
            info["city"]  = m.group(1).strip()
            info["state"] = "CA"
            info["zip"]   = m.group(2)

    return {k: v for k, v in info.items() if v and str(v).strip()}


def _cde_school_search_by_name(school_name: str, district: str = "") -> str:
    """
    Search CDE School Directory by school name (up to 100 results per page).
    Returns the URL of the best-matching detail page, or ''.
    """
    import urllib.parse
    query = urllib.parse.quote_plus(school_name)
    # Request up to 100 items to avoid missing the school on page 2
    url = (
        f"https://www.cde.ca.gov/schooldirectory/results"
        f"?searchtext={query}&searchtype=S&items=100"
    )
    html = _fetch_cde_html(url, "name search")
    if not html:
        return ""

    school_words  = [w.lower() for w in school_name.split() if len(w) > 3]
    district_word = district.lower().split()[0] if district else ""

    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, "html.parser")
        best_href   = ""
        best_score  = 0

        for a in soup.find_all("a", href=True):
            href = a["href"]
            if "details?cdscode" not in href.lower():
                continue
            link_text = a.get_text(strip=True).lower()
            score = sum(1 for w in school_words if w in link_text)
            if score == 0:
                continue
            # Bonus if district word also appears in the row
            if district_word:
                row = a.find_parent("tr")
                if row and district_word in row.get_text().lower():
                    score += 2
            if score > best_score:
                best_score = score
                best_href  = href

        if best_href:
            if not best_href.startswith("http"):
                best_href = "https://www.cde.ca.gov" + best_href
            return best_href

    except ImportError:
        m = re.search(r'href="(/[Ss]chool[Dd]irectory/details\?cdscode=\d+)"', html)
        if m:
            return "https://www.cde.ca.gov" + m.group(1)

    return ""


def _lookup_address_web(school_name: str, district: str = "", county: str = "") -> dict:
    """
    Last-resort web search (DuckDuckGo) for a school's physical address.
    """
    import urllib.parse
    context = " ".join(filter(None, [school_name, district, county, "California", "address"]))
    query   = urllib.parse.quote_plus(context)
    url     = f"https://html.duckduckgo.com/html/?q={query}"
    html    = _fetch_cde_html(url, "web address search")
    if not html:
        return {}

    m = re.search(
        r"(\d+\s+[A-Za-z0-9 ]+?"
        r"(?:Street|St|Avenue|Ave|Boulevard|Blvd|Road|Rd|Drive|Dr|Way|Lane|Ln|Court|Ct)\.?)"
        r"[,\s]+([A-Za-z ]+?),?\s*CA\s*(\d{5})",
        html, re.IGNORECASE,
    )
    if m:
        return {
            "street": m.group(1).strip(),
            "city":   m.group(2).strip().rstrip(","),
            "state":  "CA",
            "zip":    m.group(3),
        }
    return {}


def _lookup_email_web(school_name: str, district: str = "") -> str:
    """
    Search DuckDuckGo for the school's contact email address.
    """
    import urllib.parse
    query = urllib.parse.quote_plus(f"{school_name} {district} California contact email")
    html  = _fetch_cde_html(
        f"https://html.duckduckgo.com/html/?q={query}", "school email web search"
    )
    if not html:
        return ""
    emails = re.findall(r"[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}", html)
    # Prefer school/district emails over generic ones
    for e in emails:
        if "cde.ca.gov" not in e and "duckduck" not in e:
            return e
    return ""


def lookup_school_cde(cds_code: str, school_name: str = "", district: str = "") -> dict:
    """
    Look up school address and contact from the CDE School Directory.
    Strategy:
      1. CDS-code detail page
      2. Name-based CDE search  (if address still missing)
      3. DuckDuckGo web search  (final fallback for address)
      4. DuckDuckGo email search (if email still missing after above)
    """
    info = {}

    # 1. Try by CDS code
    if cds_code:
        cds_clean = normalize_cds(cds_code)
        url  = f"https://www.cde.ca.gov/schooldirectory/details?cdscode={cds_clean}"
        html = _fetch_cde_html(url, "CDS code lookup")
        if html:
            info = _parse_cde_detail_page(html)

    # 2. Name-based CDE search if address still missing
    if not info.get("street") and school_name:
        print("  [CDE] Address not found by CDS code — trying name search...")
        detail_url = _cde_school_search_by_name(school_name, district)
        if detail_url:
            html = _fetch_cde_html(detail_url, "name-search detail")
            if html:
                info = _parse_cde_detail_page(html)

    # 3. Address web fallback
    if not info.get("street") and school_name:
        print("  [Web] Searching for school address online...")
        web_info = _lookup_address_web(school_name, district)
        if web_info:
            info.update({k: v for k, v in web_info.items() if not info.get(k)})

    # 4. Email web fallback
    if not info.get("email") and school_name:
        print("  [Web] Searching for school email online...")
        email = _lookup_email_web(school_name, district)
        if email:
            info["email"] = email

    if info:
        print(f"  [CDE school] Retrieved: {', '.join(info.keys())}")
    else:
        print("  [CDE school] Address not found by any method.")
    return info


def lookup_county_from_district(district_name: str) -> str:
    """
    Find the county for a district by searching CDE School Directory.
    Falls back to a DuckDuckGo search if CDE is unreachable.
    """
    import urllib.parse

    # 1. Try CDE district search
    query = urllib.parse.quote_plus(district_name)
    url   = f"https://www.cde.ca.gov/schooldirectory/results?searchtext={query}&searchtype=D"
    html  = _fetch_cde_html(url, "district search")
    if html:
        dist_words = [w.lower() for w in district_name.split() if len(w) > 3]
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, "html.parser")
            for a in soup.find_all("a", href=True):
                href = a["href"]
                if "details?cdscode" not in href:
                    continue
                link_text = a.get_text(strip=True).lower()
                if sum(1 for w in dist_words if w in link_text) >= max(1, len(dist_words) - 1):
                    if not href.startswith("http"):
                        href = "https://www.cde.ca.gov" + href
                    detail_html = _fetch_cde_html(href, "district detail")
                    if detail_html:
                        fields = _soup_all_label_values(detail_html)
                        for k, v in fields.items():
                            if "county" in k.lower() and v:
                                return v.strip()
                    break
        except ImportError:
            pass

    # 2. Web fallback
    query2 = urllib.parse.quote_plus(f"{district_name} California county")
    html2  = _fetch_cde_html(
        f"https://html.duckduckgo.com/html/?q={query2}", "district county web search"
    )
    if html2:
        for county in CA_COUNTY_CODES.values():
            if county.lower() in html2.lower():
                return county

    return ""


def lookup_superintendent(cds_code: str) -> str:
    """
    Look up the district superintendent from the CDE School Directory district page.
    Uses the first 7 digits of the CDS code + '0000000' to get the district record.
    """
    cds_clean = normalize_cds(cds_code)
    district_cds = cds_clean[:7] + "0000000"
    url = f"https://www.cde.ca.gov/schooldirectory/details?cdscode={district_cds}"
    html = _fetch_cde_html(url, "superintendent")
    if not html:
        return ""

    fields = _soup_all_label_values(html)

    # The CDE directory lists administrators as "Administrator 1", "Administrator 2", etc.
    # alongside a title field. We look for any administrator whose title is Superintendent.
    try:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, "html.parser")
        rows = soup.find_all("tr")
        # Collect (label, value) pairs; look for administrator entries followed by title
        pairs = []
        for row in rows:
            cells = row.find_all(["th", "td"])
            if len(cells) >= 2:
                pairs.append((
                    cells[0].get_text(" ", strip=True),
                    cells[1].get_text(" ", strip=True),
                ))

        # Walk pairs: if a label contains "Administrator" and its value or the next
        # row's value contains "Superintendent", that's our person.
        for i, (lbl, val) in enumerate(pairs):
            if "administrator" in lbl.lower() and val:
                context = val.lower()
                if i + 1 < len(pairs):
                    context += " " + pairs[i + 1][1].lower()
                if "superintendent" in context:
                    name = _extract_name_only(val)
                    if name:
                        return name

        # Fallback: any field labeled with "superintendent"
        for lbl, val in pairs:
            if "superintendent" in lbl.lower() and val:
                return _extract_name_only(val)

    except ImportError:
        m = re.search(r"Superintendent[^<\n]*[:\s]+([A-Z][a-z]+ [A-Z][a-z]+)", html)
        if m:
            return m.group(1).strip()

    return ""


def _extract_name_only(text: str) -> str:
    """
    From a CDE administrator field like:
      'Dr. Kelly May-Vollmar Superintendent (760) 771-8501 kelly@school.edu'
    return just the person's name: 'Dr. Kelly May-Vollmar'
    """
    # Stop at job title keywords
    m = re.match(
        r"((?:Dr\.|Mr\.|Ms\.|Mrs\.|Prof\.)?\s*[A-Za-z][A-Za-z\-\.' ]+?)"
        r"\s+(?:Superintendent|Principal|Director|Assistant|Associate|"
        r"Administrator|Coordinator|Supervisor|Manager|Officer)",
        text,
        re.IGNORECASE,
    )
    if m:
        return m.group(1).strip()

    # Stop at phone number
    name = re.split(r"\s+\(?\d{3}\)?[\s\-\.]\d{3}", text)[0].strip()
    # Stop at email
    name = re.split(r"\s+\S+@\S+", name)[0].strip()
    return name



# Hardcoded COE Monitoring Leads table — verified from CDE contact.asp, May 2026
# Source: https://www.cde.ca.gov/ta/cr/contact.asp
_COE_LEADS_TABLE = {
    "alameda":       "Juwen Lam",
    "amador":        "Sean Snider",
    "butte":         "Susie Kruse",
    "calaveras":     "Karen Vail",
    "colusa":        "Maria Arvizu-Espinoza",
    "contra costa":  "Debra Pettric",
    "el dorado":     "Gabrielle Marchini",
    "fresno":        "Marvin Baker",
    "glenn":         "April Hine",
    "humboldt":      "August Deshais",
    "imperial":      "Claudia Montano",
    "inyo":          "Ilissa Twomey",
    "kern":          "Lily Rosenberger",
    "kings":         "Gen Almanzar",
    "lake":          "Stacie Ulatan",
    "lassen":        "James Hall",
    "los angeles":   "Adrienne Balcazar",
    "madera":        "Kirk Delmas",
    "marin":         "Laura Trahan",
    "mariposa":      "Jeff Aranguena",
    "mendocino":     "Dr. Nicole Odell",
    "merced":        "Erika Davalos-Lemus",
    "modoc":         "Mike Martin",
    "mono":          "Tammy Bennett Nguyen",
    "monterey":      "Michelle Archuleta",
    "napa":          "Lucy Pearson-Edwards",
    "nevada":        "Christine McCormick",
    "orange":        "Diane Ehrle",
    "placer":        "Leslie Wriston",
    "plumas":        "Ed Thompson",
    "riverside":     "Lisa Winberg",
    "sacramento":    "Cathy Morrison",
    "san benito":    "Mai Cruz",
    "san bernardino":"Karen Strong",
    "san diego":     "Patricia Karlin",
    "san francisco": "Mary Elisalde",
    "san joaquin":   "Sharon Oberman",
    "san luis obispo":"Stacy Summer",
    "san mateo":     "Jared Prolo",
    "santa barbara": "Shannon Yorke",
    "santa clara":   "Dawn River",
    "santa cruz":    "Angela Meeker",
    "shasta":        "Mike Freeman",
    "sierra":        "Nona Griesert",
    "siskiyou":      "Mark Lewin",
    "solano":        "Andrea Lemos",
    "sonoma":        "Amanda Welter",
    "stanislaus":    "Jill Polhemus",
    "sutter":        "Kristi Johnson",
    "tehama":        "Cathy Henderson",
    "trinity":       "Tim Nordstrom",
    "tulare":        "Gabriela Guzman",
    "tuolumne":      "Mark Pintor",
    "ventura":       "Lisa Brown",
    "yolo":          "Katrina Callaway",
    "yuba":          "Bobbi Abold",
}


def lookup_coe_lead(county_name: str) -> str:
    """
    Return the COE Monitoring Lead for a given county.

    First checks the hardcoded table (verified from CDE contact.asp, May 2026).
    Falls back to live CDE fetch if county is not in the table.
    Source: https://www.cde.ca.gov/ta/cr/contact.asp  (Compliance Monitoring > Contact Info)
    """
    import urllib.parse

    county_lower = county_name.lower().strip()

    # 1. Hardcoded table — fast, reliable, no network required
    lead = _COE_LEADS_TABLE.get(county_lower, "")
    if lead and lead != "TBD":
        return lead

    def _parse_contact_page(html: str) -> str:
        try:
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, "html.parser")
            text  = soup.get_text("\n")
        except ImportError:
            text = html

        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]

        for i, line in enumerate(lines):
            if (county_lower in line.lower()
                    and "county office of education" in line.lower()):
                for j in range(i + 1, min(i + 6, len(lines))):
                    candidate = lines[j]
                    if candidate.lower().startswith(("phone", "email", "fax", "tbd", "region")):
                        continue
                    if re.match(r"[A-Z][a-z]+(?: [A-Z][a-z]+)+", candidate):
                        return candidate
        return ""

    # 2. Live CDE fetch fallback
    html = _fetch_cde_html("https://www.cde.ca.gov/ta/cr/contact.asp", "COE leads contact")
    if html:
        lead = _parse_contact_page(html)
        if lead:
            return lead

    return ""


# ---------------------------------------------------------------------------
# Parsing the Summary of Findings
# ---------------------------------------------------------------------------

def parse_summary_of_findings(path: str):
    """
    Parse a completed CRR Summary of Findings document.

    Returns:
        metadata (dict) – school_name, cds_code, review_dates, coordinator,
                          reviewer, district, county, principal_name,
                          principal_last_name, principal_title, cde_info
        crr_sections (list of dicts)
    """
    from docx.text.paragraph import Paragraph as DocxParagraph
    from docx.table import Table as DocxTable

    doc = Document(path)

    # --- Body header fields ---
    metadata = {
        "school_name": "",
        "cds_code": "",
        "review_dates": "",
        "coordinator": "",
        "reviewer": "",
    }
    for para in doc.paragraphs[:50]:
        text = para.text.strip()
        for key, prefix in [
            ("school_name",  "School Site:"),
            ("cds_code",     "CDS Code:"),
            ("review_dates", "Review Dates:"),
            ("coordinator",  "Site CRR Coordinator:"),
            ("reviewer",     "Program Reviewer:"),
        ]:
            if text.startswith(prefix):
                metadata[key] = text.split(":", 1)[1].strip()

    # Fallback: scan table cells near the top of the document for header fields
    if any(not metadata[k] for k in ("school_name", "cds_code", "review_dates")):
        for table in doc.tables[:5]:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    for key, prefix in [
                        ("school_name",  "School Site:"),
                        ("cds_code",     "CDS Code:"),
                        ("review_dates", "Review Dates:"),
                        ("coordinator",  "Site CRR Coordinator:"),
                        ("reviewer",     "Program Reviewer:"),
                    ]:
                        if not metadata[key] and text.startswith(prefix):
                            metadata[key] = text.split(":", 1)[1].strip()

    # --- District from page header paragraphs ---
    district = ""
    for section in doc.sections:
        for p in section.header.paragraphs:
            t = p.text.strip()
            if t and t != metadata.get("school_name", "") and any(
                kw in t for kw in ["District", "Unified", "Elementary", "High", "Charter"]
            ):
                district = t
                break
        if district:
            break
    metadata["district"] = district

    # --- County from CDS code ---
    metadata["county"] = county_from_cds(metadata.get("cds_code", ""))

    # --- Principal from coordinator field ---
    coord = metadata.get("coordinator", "")
    if coord:
        principal_name, principal_last, principal_title = parse_principal(coord)
    else:
        principal_name = principal_last = principal_title = ""
    metadata["principal_name"]       = principal_name
    metadata["principal_last_name"]  = principal_last
    metadata["principal_title"]      = principal_title

    # --- CDE web lookups ---
    cds      = metadata.get("cds_code", "")
    school   = metadata.get("school_name", "")
    district = metadata.get("district", "")
    county   = metadata.get("county", "")

    print("  Looking up school info from CDE School Directory...")
    metadata["cde_info"] = lookup_school_cde(cds, school, district)

    # --- County fallback: use district name if CDS code gave nothing ---
    if not county and district:
        print("  County not found from CDS code — looking up via district name...")
        county = lookup_county_from_district(district)
        if county:
            metadata["county"] = county
            print(f"  [County] Found via district: {county}")

    print("  Looking up superintendent from CDE district page...")
    metadata["superintendent"] = lookup_superintendent(cds)
    if metadata["superintendent"]:
        print(f"  [CDE district] Superintendent: {metadata['superintendent']}")

    print("  Looking up COE Monitoring Lead from CDE Compliance Monitoring page...")
    metadata["coe_lead"] = lookup_coe_lead(county)
    if metadata["coe_lead"]:
        print(f"  [CDE leads] COE Lead ({county}): {metadata['coe_lead']}")

    # --- Walk body elements for CRR sections ---
    crr_sections = []
    current_crr = None
    state = None

    for element in doc.element.body:
        tag = element.tag.split("}")[-1]

        if tag == "p":
            para = DocxParagraph(element, doc)
            text = para.text.strip()
            style = para.style.name if para.style else ""

            if style == "Heading 1":
                m = re.match(r"(CRR\s+\d+):\s*(.+)", text)
                if m:
                    if current_crr is not None:
                        crr_sections.append(current_crr)
                    crr_num = re.sub(r"\s+", " ", m.group(1)).strip()
                    current_crr = {
                        "crr_num": crr_num,
                        "crr_title": m.group(2).strip(),
                        "corrective_actions": [],
                        "is_accessible_facilities": "accessible facilit" in text.lower(),
                    }
                    state = "heading"
                elif current_crr is not None and "REMOVED" in text:
                    crr_sections.append(current_crr)
                    current_crr = None
                    state = None

            elif style == "Heading 2":
                if "Required Corrective Action" in text:
                    state = "corrective_actions"
                elif "Summary of" in text or "Analysis" in text:
                    state = "summary"
                elif "Observation" in text:
                    state = "observation"
                else:
                    state = "other"

            elif state == "corrective_actions" and current_crr:
                clean = text.replace(" ", "").strip()
                if clean:
                    current_crr["corrective_actions"].append(clean)

        elif tag == "tbl":
            if current_crr and current_crr.get("is_accessible_facilities"):
                table = DocxTable(element, doc)
                for row in table.rows[1:]:
                    cells = row.cells
                    if not cells:
                        continue
                    area       = cells[0].text.strip()
                    correction = cells[-1].text.strip()
                    if correction and not is_none_action(correction):
                        current_crr["corrective_actions"].append(
                            f"{area}\nCorrective Action: {correction}"
                        )

    if current_crr is not None:
        crr_sections.append(current_crr)

    return metadata, crr_sections


def get_findings(crr_sections: list) -> list:
    findings = []
    for crr in crr_sections:
        real = [a for a in crr["corrective_actions"] if not is_none_action(a)]
        if real:
            findings.append({**crr, "corrective_actions": real})
    return findings


# ---------------------------------------------------------------------------
# Text replacement utilities
# ---------------------------------------------------------------------------

def _replace_adjacent_runs(para, old: str, new: str, max_window: int = 6) -> bool:
    """
    Replace `old` text that may be split across up to max_window consecutive runs.
    Only modifies the minimum number of runs needed.
    """
    runs = para.runs
    n = len(runs)
    for size in range(1, min(max_window + 1, n + 1)):
        for i in range(n - size + 1):
            combined = "".join(r.text for r in runs[i:i + size])
            if old in combined:
                runs[i].text = combined.replace(old, new, 1)
                for j in range(i + 1, i + size):
                    runs[j].text = ""
                return True
    return False


def replace_in_doc(doc, replacements: dict):
    """Apply replacement dict across all paragraphs (body + headers + footers)."""
    all_para_sets = [doc.paragraphs]
    for section in doc.sections:
        for attr in ("header", "footer"):
            try:
                all_para_sets.append(getattr(section, attr).paragraphs)
            except Exception:
                pass

    for paras in all_para_sets:
        for para in paras:
            for old, new in replacements.items():
                # Loop to replace every occurrence, not just the first
                while _replace_adjacent_runs(para, old, new):
                    pass


# ---------------------------------------------------------------------------
# VCP generation
# ---------------------------------------------------------------------------

def _set_cell_content(cell, title: str, body_lines: list):
    for para in cell.paragraphs:
        para.clear()
    run = cell.paragraphs[0].add_run(title)
    run.bold = True
    for line in body_lines:
        for subline in str(line).split("\n"):
            cell.add_paragraph(subline)


def fill_vcp(template_path: str, output_path: str, metadata: dict, findings: list, today: datetime):
    doc = Document(template_path)

    school_name  = metadata.get("school_name") or "[School Name]"
    review_dates = metadata.get("review_dates") or "[Date]"
    cds_code     = metadata.get("cds_code") or "[CDS Code]"

    # 45-day deadline is from the LOF cover letter date (today), not review end date
    deadline_str = (today + timedelta(days=45)).strftime("%B %d, %Y")

    replace_in_doc(doc, {
        "[School Name]":        school_name,
        "[Date]":               review_dates,
        "[45 days after Date]": deadline_str,
        "[CDS Code]":           cds_code,
        # Also handle split-run variants
        "School Name":          school_name,
        "45 days after Date":   deadline_str,
        "CDS Code":             cds_code,
    })

    # Fill the table
    table = doc.tables[0]
    blank_trs = [
        row._tr for row in table.rows[1:]
        if not any(cell.text.strip() for cell in row.cells)
    ]
    for tr in blank_trs:
        tr.getparent().remove(tr)

    if not findings:
        new_row = table.add_row()
        new_row.cells[0].text = "N/A"
        new_row.cells[1].text = "No findings of noncompliance were identified."
        for i in range(2, min(6, len(new_row.cells))):
            new_row.cells[i].text = "N/A"
    else:
        for finding in findings:
            new_row = table.add_row()
            new_row.cells[0].text = finding["crr_num"]
            _set_cell_content(
                new_row.cells[1],
                title=f"{finding['crr_num']}: {finding['crr_title']}",
                body_lines=finding["corrective_actions"],
            )

    doc.save(output_path)
    print(f"  [OK] VCP saved: {output_path}")


# ---------------------------------------------------------------------------
# LOF Cover Letter generation
# ---------------------------------------------------------------------------

def _fill_address_block(doc, metadata: dict, today: datetime):
    """Fill in the address block, salutation, and cc list of the cover letter."""
    school_name        = metadata.get("school_name", "")
    principal_name     = metadata.get("principal_name", "")
    principal_last     = metadata.get("principal_last_name", "")
    principal_title    = metadata.get("principal_title", "Principal")
    coordinator_name   = metadata.get("principal_name", "")  # coordinator IS the principal
    district           = metadata.get("district", "")
    county             = metadata.get("county", "")
    superintendent     = metadata.get("superintendent", "")
    coe_lead           = metadata.get("coe_lead", "")
    cde                = metadata.get("cde_info", {})

    street  = cde.get("street", "")
    city    = cde.get("city", "")
    state   = cde.get("state", "CA")
    zip_    = cde.get("zip", "")
    email   = cde.get("email", "")

    city_state_zip = f"{city}, {state} {zip_}".strip(", ") if city else ""

    paras = doc.paragraphs

    # Para [4]: "Name, Principal"
    if principal_name:
        title_str = principal_title if principal_title else "Principal"
        _replace_adjacent_runs(paras[4], "Name, Principal", f"{principal_name}, {title_str}")

    # Para [5]: "School Site Name"
    if school_name:
        _replace_adjacent_runs(paras[5], "School Site Name", school_name)
        _replace_adjacent_runs(paras[5], "School Site Name ", school_name)

    # Para [6]: "Address"
    if street:
        for run in paras[6].runs:
            if run.text.strip() == "Address":
                run.text = street
                break

    # Para [7]: "City, State Zip Code"
    p7 = paras[7]
    if city_state_zip:
        _replace_adjacent_runs(p7, "City, State Zip Code", city_state_zip)

    if principal_last:
        _replace_adjacent_runs(p7, "[Last name]", principal_last)

    # Email: search ALL paragraphs for the placeholder (template placement varies)
    EMAIL_PLACEHOLDERS = [
        "Email Address < When finalizing letter, this should be a hyperlink (in blue) ",
        "Email Address",
        "[Email]",
        "[Email Address]",
    ]
    email_placed = False
    for para in paras:
        for placeholder in EMAIL_PLACEHOLDERS:
            if email:
                if _replace_adjacent_runs(para, placeholder, email, max_window=10):
                    email_placed = True
                    break
            else:
                # Remove placeholder so it doesn't appear in final letter
                if _replace_adjacent_runs(para, placeholder, "", max_window=10):
                    break
        if email_placed:
            break

    # If no placeholder found but we have an email, insert it inside p7
    # right after the first <w:br/> (which immediately follows the city/zip text)
    # so it appears between the city/zip line and the "Dear Principal" salutation.
    if email and not email_placed:
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        p7_runs = p7._p.findall(qn("w:r"))
        first_br_run = next(
            (r for r in p7_runs if r.find(qn("w:br")) is not None), None
        )
        if first_br_run is not None:
            email_run = OxmlElement("w:r")
            email_t   = OxmlElement("w:t")
            email_t.text = email
            email_run.append(email_t)
            first_br_run.addnext(email_run)
        else:
            # Fallback: insert as a new paragraph after p7
            from copy import deepcopy
            new_para = deepcopy(p7._p)
            for r in new_para.findall(qn("w:r")):
                new_para.remove(r)
            r_elem = OxmlElement("w:r")
            t_elem = OxmlElement("w:t")
            t_elem.text = email
            r_elem.append(t_elem)
            new_para.append(r_elem)
            p7._p.addnext(new_para)

    # Para [25]: cc list
    p25 = paras[25]
    if coordinator_name:
        _replace_adjacent_runs(p25, "[Name], Designated CRR Coordinator",
                               f"{coordinator_name}, Designated CRR Coordinator")
    if school_name:
        _replace_adjacent_runs(p25, "[School Site]", school_name)
        _replace_adjacent_runs(p25, "[School Site] ", school_name + " ")
    if superintendent:
        _replace_adjacent_runs(p25, "[Name], Superintendent",
                               f"{superintendent}, Superintendent")
    if district:
        _replace_adjacent_runs(p25, "[School District]", district)
        _replace_adjacent_runs(p25, "[School District] ", district + " ")
    if coe_lead:
        _replace_adjacent_runs(p25, "[Name], COE Monitoring Lead",
                               f"{coe_lead}, COE Monitoring Lead")
    if county:
        _replace_adjacent_runs(p25, "[County]", county)
        _replace_adjacent_runs(p25, "[County] ", county + " ")


def fill_cover_letter(
    template_path: str,
    output_path: str,
    metadata: dict,
    has_findings: bool,
    today: datetime,
    preparer_initials: str = "MM",
):
    doc = Document(template_path)

    school_name  = metadata.get("school_name") or "[School Site Name]"
    review_dates = metadata.get("review_dates") or "[dates]"
    deadline_str = (today + timedelta(days=45)).strftime("%B %d, %Y")
    today_str    = today.strftime("%B %d, %Y")

    # Signature: "RST:[preparer initials]" — RST = Randi Solís Thompson (Civil Rights Officer)
    # preparer_initials = initials of the OEO staff member who prepared the letter (default MM)

    # Global text replacements
    replace_in_doc(doc, {
        "[Date]":              today_str,
        "[School Site Name]":  school_name,
        "[School Site]":       school_name,
        "[dates]":             review_dates,
        "[45 calendar days]":  deadline_str,
        "[initials]":          preparer_initials,
    })

    # Address block, salutation, cc list
    _fill_address_block(doc, metadata, today)

    doc.save(output_path)
    letter_type = "Findings" if has_findings else "No Findings"
    print(f"  [OK] LOF Cover Letter ({letter_type}) saved: {output_path}")


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="CRR Civil Rights Review Document Processing Agent",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("input", help="Path to the completed CRR Summary of Findings (.docx)")
    parser.add_argument("--output-dir", "-o", default="outputs",
                        help="Output directory (default: outputs/)")
    parser.add_argument("--vcp-template",
                        default=str(TEMPLATE_DIR / "VCP_template.docx"))
    parser.add_argument("--lof-findings-template",
                        default=str(TEMPLATE_DIR / "LOF_Findings_template.docx"))
    parser.add_argument("--lof-no-findings-template",
                        default=str(TEMPLATE_DIR / "LOF_NoFindings_template.docx"))
    parser.add_argument("--preparer", default="MM",
                        help="Initials of the OEO staff member preparing the letter "
                             "(used in 'RST:[initials]' signature line). Default: MM")

    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f"ERROR: File not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    today = datetime.now()

    print()
    print("=" * 62)
    print("  CRR Document Processing Agent")
    print("=" * 62)
    print(f"  Input : {input_path.name}")
    print()

    metadata, crr_sections = parse_summary_of_findings(str(input_path))
    findings    = get_findings(crr_sections)
    has_findings = bool(findings)

    print()
    print(f"  School        : {metadata.get('school_name') or '(not found)'}")
    print(f"  District      : {metadata.get('district') or '(not found)'}")
    print(f"  County        : {metadata.get('county') or '(not found)'}")
    print(f"  Principal     : {metadata.get('principal_name') or '(not found)'}")
    print(f"  Superintendent: {metadata.get('superintendent') or '(not found - blank in letter)'}")
    print(f"  COE Lead      : {metadata.get('coe_lead') or '(not found - blank in letter)'}")
    print(f"  Review Dates  : {metadata.get('review_dates') or '(not found)'}")
    print(f"  CRR Sections : {len(crr_sections)} parsed")
    print(f"  Findings     : {len(findings)} section(s) with corrective action(s)")

    if findings:
        print()
        print("  Sections with findings:")
        for f in findings:
            print(f"    * {f['crr_num']}: {f['crr_title']}")
    else:
        print("  -> No findings. Generating 'No Findings' cover letter.")

    print()
    print(f"  LOF Date     : {today.strftime('%B %d, %Y')}")
    print(f"  VCP Deadline : {(today + timedelta(days=45)).strftime('%B %d, %Y')} (45 days from LOF date)")
    print(f"  Cover Letter : {'LOF - Findings' if has_findings else 'LOF - No Findings'}")
    print()

    safe = re.sub(r"[^\w\s\-]", "", metadata.get("school_name", "School"))
    safe = safe.strip().replace(" ", "_")

    vcp_out = output_dir / f"{safe}_VCP.docx"
    lof_out = output_dir / f"{safe}_LOF_Cover_Letter.docx"

    # Detect a naming conflict only when a file with this name was created very recently
    # (within 10 minutes) — meaning another school in the same batch already used this name.
    # Re-runs on a later day will simply overwrite the old file (same school, same name).
    import time as _time
    _now = _time.time()
    _recent = 600  # seconds
    _conflict = (
        (vcp_out.exists() and (_now - vcp_out.stat().st_mtime) < _recent) or
        (lof_out.exists() and (_now - lof_out.stat().st_mtime) < _recent)
    )
    if _conflict:
        district_raw = metadata.get("district", "")
        # Build a meaningful suffix from the first 1-2 distinctive district words
        _skip = {"unified", "high", "school", "district", "county",
                 "the", "of", "and", "valley", "union", "joint"}
        dist_words = [w for w in district_raw.split() if w.lower() not in _skip]
        # Use up to 2 words for a readable suffix (e.g. "San_Ramon" not just "San")
        suffix = "_".join(dist_words[:2]) if dist_words else "Alt"
        safe   = f"{safe}_{suffix}"
        vcp_out = output_dir / f"{safe}_VCP.docx"
        lof_out = output_dir / f"{safe}_LOF_Cover_Letter.docx"
        print(f"  [Note] Name conflict — using '{safe}' as filename.")

    fill_vcp(args.vcp_template, str(vcp_out), metadata, findings, today)

    lof_template = (
        args.lof_findings_template if has_findings else args.lof_no_findings_template
    )
    fill_cover_letter(lof_template, str(lof_out), metadata, has_findings, today,
                      preparer_initials=args.preparer)

    print()
    print("  Generated files:")
    print(f"    {vcp_out}")
    print(f"    {lof_out}")
    print("=" * 62)
    print()


if __name__ == "__main__":
    main()
