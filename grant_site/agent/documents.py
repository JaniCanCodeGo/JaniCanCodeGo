"""Render structured documents to Markdown and (when python-docx is
installed) to .docx files in the run's output directory."""

import json
import os
import re

try:
    from docx import Document
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


def safe_name(text):
    text = re.sub(r"[^\w\s-]", "", text or "document").strip()
    return re.sub(r"\s+", "_", text)[:80] or "document"


def render_markdown(doc):
    lines = [f"# {doc['title']}", ""]
    for heading, body in doc["sections"]:
        lines += [f"## {heading}", "", body, ""]
    return "\n".join(lines)


def write_document(doc, out_dir, basename):
    """Write markdown (always) and docx (if available). Returns file list."""
    os.makedirs(out_dir, exist_ok=True)
    files = []
    md_path = os.path.join(out_dir, f"{basename}.md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(render_markdown(doc))
    files.append(md_path)

    if HAS_DOCX:
        word = Document()
        word.add_heading(doc["title"], level=0)
        for heading, body in doc["sections"]:
            word.add_heading(heading, level=1)
            for para in body.split("\n\n"):
                word.add_paragraph(para)
        docx_path = os.path.join(out_dir, f"{basename}.docx")
        word.save(docx_path)
        files.append(docx_path)
    return files


def write_json(data, out_dir, basename):
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, f"{basename}.json")
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
    return [path]
