#!/usr/bin/env python3
"""CETL Business Plan Agent.

Two jobs, one agent:

1. `write` — pull together everything this repo knows about CETL (the site
   manifest, curriculum, podcast episodes, README, and the verified research
   footnotes already in BUSINESS_PLAN.md) and build a complete
   write-the-business-plan prompt. No API key needed — paste the prompt into
   Claude, save Claude's markdown output over BUSINESS_PLAN.md, then run
   the docx build below. (Also exposed as an MCP tool via cetl_mcp_server.py.)

2. `docx` — convert BUSINESS_PLAN.md into a formatted, downloadable Word
   document at outputs/CETL_Business_Plan.docx.

Usage:
    python3 business_plan_agent.py                 # builds the .docx (default)
    python3 business_plan_agent.py docx
    python3 business_plan_agent.py docx --source BUSINESS_PLAN.md --output-dir outputs
    python3 business_plan_agent.py write           # builds the full writing prompt
    python3 business_plan_agent.py write --focus "update the financials for year 2"

The markdown file stays the source of truth; the .docx is a build artifact
(outputs/ is gitignored).

Requires: python-docx  (pip install python-docx)
"""

import argparse
import json
import os
import re
import sys
from datetime import datetime

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("[Error] python-docx is required: pip install python-docx")
    sys.exit(1)

ACCENT = RGBColor(0x5B, 0x2D, 0x86)  # deep plum — warm, not corporate-gray
BODY_FONT = "Calibri"


def _add_runs_with_inline_formatting(paragraph, text):
    """Render a markdown text fragment into runs: **bold**, *italic*,
    [text](url) links (as 'text (url)'), [^n] footnote refs (superscript),
    and `code` (monospace)."""
    # normalize links first so their brackets don't confuse other patterns
    text = re.sub(r"\[([^\]^][^\]]*)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    token_re = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`|\[\^\d+\])")
    for part in token_re.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif re.fullmatch(r"\[\^\d+\]", part):
            run = paragraph.add_run(part[2:-1])
            run.font.superscript = True
        else:
            paragraph.add_run(part)


def _add_table(doc, rows):
    """rows: list of lists of cell strings; first row is the header."""
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.cell(r, c)
            cell.paragraphs[0].text = ""
            _add_runs_with_inline_formatting(cell.paragraphs[0], cell_text)
            if r == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def _style_document(doc):
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(11)
    for level, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        h = doc.styles[level]
        h.font.name = BODY_FONT
        h.font.size = Pt(size)
        h.font.color.rgb = ACCENT
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)


def _split_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def build_docx(source_path, output_path):
    with open(source_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    _style_document(doc)

    footnotes = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        # footnote definitions are collected and rendered as an endnotes section
        fn = re.match(r"\[\^(\d+)\]:\s*(.*)", stripped)
        if fn:
            footnotes.append((fn.group(1), fn.group(2)))
            i += 1
            continue

        heading = re.match(r"(#{1,4})\s+(.*)", stripped)
        if heading:
            if heading.group(2).strip() == "Footnotes":
                # rendered as its own endnotes section after the body
                i += 1
                continue
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                p = doc.add_heading("", level=0)
                _add_runs_with_inline_formatting(p, text)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p = doc.add_heading("", level=min(level - 1, 3))
                _add_runs_with_inline_formatting(p, text)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = _split_table_row(lines[i])
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                _add_table(doc, rows)
            continue

        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph(style="Intense Quote")
            _add_runs_with_inline_formatting(p, " ".join(q for q in quote if q))
            continue

        bullet = re.match(r"[-*]\s+(.*)", stripped)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_inline_formatting(p, bullet.group(1))
            i += 1
            continue

        numbered = re.match(r"\d+\.\s+(.*)", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            _add_runs_with_inline_formatting(p, numbered.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        _add_runs_with_inline_formatting(p, stripped)
        i += 1

    if footnotes:
        doc.add_page_break()
        doc.add_heading("Footnotes", level=1)
        for num, text in footnotes:
            p = doc.add_paragraph()
            run = p.add_run(num)
            run.font.superscript = True
            p.add_run("  ")
            _add_runs_with_inline_formatting(p, text)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)
    return output_path


# ---------------------------------------------------------------------------
# `write` — gather everything the repo knows and build the writing prompt
# ---------------------------------------------------------------------------

HERE = os.path.dirname(os.path.abspath(__file__))


def _read(path, limit=None):
    try:
        with open(path, encoding="utf-8") as f:
            text = f.read()
        return text[:limit] + "\n[...truncated...]" if limit and len(text) > limit else text
    except OSError:
        return None


def _extract_research(plan_path):
    """Pull the verified research out of an existing BUSINESS_PLAN.md so a
    rewrite keeps its fact-checked sources instead of inventing new ones."""
    text = _read(plan_path)
    if not text:
        return None
    footnotes = re.findall(r"^\[\^\d+\]:.*$", text, flags=re.MULTILINE)
    return "\n".join(footnotes) if footnotes else None


def gather_project_information():
    """Collect everything in this repo that should inform the business plan."""
    sections = []

    manifest = _read(os.path.join(HERE, "data", "site_manifest.json"))
    if manifest:
        sections.append(("Project manifest (data/site_manifest.json)", manifest))

    curriculum = _read(os.path.join(HERE, "data", "curriculum.json"))
    if curriculum:
        sections.append(("Curriculum — the Build With Claude product (data/curriculum.json)", curriculum))

    episodes = _read(os.path.join(HERE, "data", "podcast_episodes.json"))
    if episodes:
        sections.append(("Podcast episode list (data/podcast_episodes.json)", episodes))

    readme = _read(os.path.join(HERE, "README.md"))
    if readme:
        sections.append(("Project README", readme))

    research = _extract_research(os.path.join(HERE, "BUSINESS_PLAN.md"))
    if research:
        sections.append((
            "Verified research footnotes from the current business plan "
            "(each was independently fact-checked — reuse these; do not invent new statistics)",
            research,
        ))

    promo_dir = os.path.join(HERE, "content", "promo")
    if os.path.isdir(promo_dir):
        for name in sorted(os.listdir(promo_dir))[:3]:
            text = _read(os.path.join(promo_dir, name), limit=2500)
            if text:
                sections.append((f"Marketing voice sample (content/promo/{name})", text))

    return sections


BUSINESS_PLAN_REQUEST = """Write a complete business plan for Clear Enough To Lead (CETL) in GitHub-flavored markdown, using ONLY the project information and verified research provided above.

Structure (keep these sections, in this order):
1. Executive Summary
2. Company Overview (mission, what CETL is not, founder, legal/structure)
3. The Problem — Researched Pain Points
4. Market Analysis (top-down and bottom-up, plus ideal customer profile)
5. The Solution — and Why It's Evidence-Based (map each product design choice to its evidence)
6. Competitive Landscape (table)
7. Marketing & Customer Acquisition
8. Revenue Model & Financial Projections (label projections as founder assumptions)
9. Operations & Roadmap (phases and KPIs)
10. Risks & Mitigations (table)
Footnotes (numbered markdown footnotes with URLs)

Hard rules:
- Every statistic or market claim must cite one of the provided verified footnotes using [^n] markers. If a claim has no provided source, either omit it or clearly label it as a founder assumption.
- Do NOT invent statistics, market sizes, or study results beyond the provided research.
- Financial projections are planning assumptions — say so explicitly.
- Voice: professional business-plan register, but plainspoken and honest in the CETL spirit — no hype, no buzzwords, no toxic positivity. Flag weak evidence honestly.
- The output must be pure markdown, starting with a single `# ` title line, compatible with this agent's docx builder (headings, tables, bullet lists, blockquotes, and [^n] footnotes only).

Additional focus for this revision: {focus}

After generating, save the markdown over CETL/BUSINESS_PLAN.md and run `python3 business_plan_agent.py` to produce the Word document."""


def build_business_plan_prompt(focus: str = "") -> str:
    parts = ["# CETL project information (gathered automatically)\n"]
    for title, body in gather_project_information():
        parts.append(f"## {title}\n\n{body.strip()}\n")
    parts.append("---\n")
    parts.append(BUSINESS_PLAN_REQUEST.format(focus=focus.strip() or "none — full plan refresh"))
    return "\n".join(parts)


def save_prompt(prompt: str) -> str:
    folder = os.path.join(HERE, "outputs", "business_plan")
    os.makedirs(folder, exist_ok=True)
    path = os.path.join(folder, f"write_prompt_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(prompt)
    return path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="CETL business plan agent: gather info + build the writing prompt, or build the .docx.",
    )
    sub = parser.add_subparsers(dest="command")

    docx_cmd = sub.add_parser("docx", help="Convert BUSINESS_PLAN.md to a Word document (default)")
    docx_cmd.add_argument("--source", default=os.path.join(HERE, "BUSINESS_PLAN.md"))
    docx_cmd.add_argument("--output-dir", default=os.path.join(HERE, "outputs"))

    write_cmd = sub.add_parser("write", help="Gather all project info and build the plan-writing prompt")
    write_cmd.add_argument("--focus", "-f", default="", help="What this revision should emphasize or update")
    write_cmd.add_argument("--no-save", action="store_true")

    # keep the original no-argument behavior: build the docx
    args = parser.parse_args()
    command = args.command or "docx"

    if command == "write":
        prompt = build_business_plan_prompt(focus=args.focus)
        print(prompt)
        if not args.no_save:
            path = save_prompt(prompt)
            print(f"\n[OK] Prompt saved to {path}")
            print("     -> Paste it into Claude, save the result over BUSINESS_PLAN.md,")
            print("        then run: python3 business_plan_agent.py")
        return

    source = getattr(args, "source", os.path.join(HERE, "BUSINESS_PLAN.md"))
    output_dir = getattr(args, "output_dir", os.path.join(HERE, "outputs"))
    if not os.path.exists(source):
        print(f"[Error] Source not found: {source}")
        sys.exit(1)
    out = build_docx(source, os.path.join(output_dir, "CETL_Business_Plan.docx"))
    print(f"[OK] Business plan written to {out}")


if __name__ == "__main__":
    main()
